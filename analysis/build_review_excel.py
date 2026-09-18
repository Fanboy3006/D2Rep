#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_review_excel.py - 把 Q2/Q3 分析结果整合成"方便人类复核"的 xlsx。

风格参考 owner 的 XG/VG/TS统计.xlsx：Sheet1=逐场明细，Sheet2=汇总/分布，
外加口径说明。产出:
  analysis/output_review/DOTA_review.xlsx

数据来源:
  - stats.db  (Q2: matches + gold_adv + team 名)  -> 逐场明细 + 每队汇总
  - analysis/output_q3/q3a_hero_nw_increment.csv   -> Q3a 英雄每分净值增量
  - analysis/output_q3/q3a_rel_team_networth.csv   -> Q3a 相对量
  - analysis/output_q3/q3b_hero_state_winrate.csv  -> Q3b 经济->胜率
用法:
  python analysis/build_review_excel.py
"""
import csv
import os
import sqlite3
import sys

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STATS = os.path.join(ROOT, "stats.db")
Q3 = os.path.join(ROOT, "analysis", "output_q3")
OUT = os.path.join(ROOT, "analysis", "output_review")
os.makedirs(OUT, exist_ok=True)
XLSX = os.path.join(OUT, "DOTA_review.xlsx")
THR = 1000
TEAM_NAME = {2: "Radiant", 3: "Dire"}


def read_csv(path):
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def main():
    con = sqlite3.connect(STATS)
    con.row_factory = sqlite3.Row
    wb = openpyxl.Workbook()

    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    win_fill = PatternFill("solid", fgColor="C6EFCE")
    lose_fill = PatternFill("solid", fgColor="FFC7CE")

    def style_header(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center")

    # ---------- 口径 sheet ----------
    ws0 = wb.active
    ws0.title = "0_说明口径"
    lines = [
        "DOTA 分析 · 人类复核用表",
        "",
        "口径：",
        "  A1 逐玩家逐分钟净值 = 读 .dem 的 CDOTA_DataRadiant/Dire.m_iNetWorth（与 OpenDota 对账 0.000%）。",
        "  队伍净值(min) = 该队5人逐秒净值求和；gold_adv(min) = Radiant净值 - Dire净值。",
        "  10min状态: >+%d = lead(领先)  <-%d = trail(落后)  之间=even(均势)。" % (THR, THR),
        "  Δ = gold_adv[20] - gold_adv[10]; >+%d=turn_up  <-%d=turn_down  之间=flat。" % (THR, THR),
        "  team_win = radiant_win if 该队radiant else 1-radiant_win。",
        "  Q3a_rel: team_nw_gain/min=英雄在队时队伍每分净值增量; delta_adv/min=双方净值差每分变化(该队视角), 正=IMPROVE 负=DECREASE。",
        "  0-10窗口因部分录像'中途起录'偏空; n 为样本数, 多数较小属方向性。",
        "",
        "文件: " + XLSX,
    ]
    for i, l in enumerate(lines, 1):
        ws0.cell(row=i, column=1, value=l)
    ws0.column_dimensions["A"].width = 120

    # ---------- Q2 逐场明细 ----------
    rows = []
    # per match, per team in {XG,VG,TS,TY,PV} from matches table (we reuse stats.db which has these)
    target_tids = {8261500: "XG", 726228: "VG", 7119388: "TS", 9823272: "TY",
                   9572001: "PV", 9824702: "PV"}
    q = con.execute("""SELECT m.match_id,m.radiant_team_id,m.dire_team_id,m.radiant_win,
                       (SELECT name FROM teams t WHERE t.team_id=m.radiant_team_id) rn,
                       (SELECT name FROM teams t WHERE t.team_id=m.dire_team_id) dn
                       FROM matches m""")
    for m in q:
        mid = m["match_id"]
        g10r = con.execute("SELECT value FROM gold_adv WHERE match_id=? AND minute=10", (mid,)).fetchone()
        g20r = con.execute("SELECT value FROM gold_adv WHERE match_id=? AND minute=20", (mid,)).fetchone()
        if not g10r or not g20r:
            continue
        g10 = g10r["value"]; g20 = g20r["value"]
        for side, tid, team, win in (("Radiant", m["radiant_team_id"], m["rn"], m["radiant_win"]),
                                      ("Dire", m["dire_team_id"], m["dn"], m["radiant_win"])):
            label = target_tids.get(tid)
            if label is None:
                continue
            if side == "Radiant":
                t10, t20, res = g10, g20, win
            else:
                t10, t20 = -g10, -g20
                res = 1 - win if win is not None else None
            dw = t20 - t10
            s10 = "lead" if t10 > THR else ("trail" if t10 < -THR else "even")
            ds = "turn_up" if dw > THR else ("turn_down" if dw < -THR else "flat")
            rows.append({"team": label, "match_id": mid, "side": side, "opp": (m["dn"] if side == "Radiant" else m["rn"]),
                         "g10": t10, "g20": t20, "delta": dw, "res": ("W" if res == 1 else ("L" if res == 0 else "?")),
                         "s10": s10, "s_change": ds})
    ws1 = wb.create_sheet("1_Q2_逐场明细")
    hdr = ["team", "match_id", "side", "opp", "10min_goldadv", "20min_goldadv", "delta(20-10)", "result", "10min_state", "10to20"]
    ws1.append(hdr); style_header(ws1, 1, len(hdr))
    for r in sorted(rows, key=lambda x: (x["team"], x["match_id"])):
        ws1.append([r["team"], r["match_id"], r["side"], r["opp"], r["g10"], r["g20"], r["delta"], r["res"], r["s10"], r["s_change"]])
    for i in range(2, ws1.max_row + 1):
        if ws1.cell(row=i, column=8).value == "W":
            ws1.cell(row=i, column=8).fill = win_fill
        elif ws1.cell(row=i, column=8).value == "L":
            ws1.cell(row=i, column=8).fill = lose_fill
    for col, wd in zip("ABCDEFGHIJ", [8, 14, 8, 12, 14, 14, 14, 8, 12, 12]):
        ws1.column_dimensions[col].width = wd

    # ---------- Q2 每队汇总 ----------
    ws2 = wb.create_sheet("2_Q2_队伍汇总")
    ws2.append(["team", "matches", "lead_wr", "even_wr", "trail_wr", "turn_up_wr", "flat_wr", "turn_down_wr"])
    style_header(ws2, 1, 8)
    import collections
    per = collections.defaultdict(lambda: collections.defaultdict(lambda: [0, 0]))
    for r in rows:
        p = per[r["team"]]
        for k in ("lead", "even", "trail"):
            if r["s10"] == k:
                p[k][0] += 1; p[k][1] += 1 if r["res"] == "W" else 0
        for k in ("turn_up", "flat", "turn_down"):
            if r["s_change"] == k:
                p[k][0] += 1; p[k][1] += 1 if r["res"] == "W" else 0
    for team in sorted(per):
        p = per[team]
        def wr(k):
            n, w = p[k]; return round(w / n, 4) if n else None
        ws2.append([team, len([r for r in rows if r["team"] == team]),
                    wr("lead"), wr("even"), wr("trail"),
                    wr("turn_up"), wr("flat"), wr("turn_down")])
    for col in "ABCDEFGH":
        ws2.column_dimensions[col].width = 12

    # ---------- Q3a 英雄每分净值增量 ----------
    q3a = read_csv(os.path.join(Q3, "q3a_hero_nw_increment.csv"))
    if q3a:
        ws3 = wb.create_sheet("3_Q3a_英雄净值增量")
        ws3.append(["hero", "0-10/min", "10-20/min", "20+/min", "n_0_10", "n_10_20", "n_20"])
        style_header(ws3, 1, 7)
        for r in q3a:
            ws3.append([r.get("hero"), r.get("per_min_0_10"), r.get("per_min_10_20"),
                        r.get("per_min_20_plus"), r.get("n0_10"), r.get("n10_20"), r.get("n20")])
        for col in "ABCDEFG":
            ws3.column_dimensions[col].width = 16

    # ---------- Q3a 相对量 ----------
    rel = read_csv(os.path.join(Q3, "q3a_rel_team_networth.csv"))
    if rel:
        ws4 = wb.create_sheet("4_Q3a_队伍净值与优劣")
        ws4.append(["hero", "window", "team_nw_gain_per_min", "delta_adv_per_min", "net_view", "n"])
        style_header(ws4, 1, 6)
        for r in rel:
            ws4.append([r.get("hero"), r.get("window"), r.get("team_nw_gain_per_min"),
                        r.get("delta_adv_per_min"), r.get("net_view"), r.get("n")])
        for col in "ABCDEF":
            ws4.column_dimensions[col].width = 18

    # ---------- Q3b 经济->胜率 ----------
    q3b = read_csv(os.path.join(Q3, "q3b_hero_state_winrate.csv"))
    if q3b:
        ws5 = wb.create_sheet("5_Q3b_经济到胜率")
        ws5.append(["hero", "state", "n", "win", "winrate"])
        style_header(ws5, 1, 5)
        for r in q3b:
            ws5.append([r.get("hero"), r.get("state"), r.get("n"), r.get("win"), r.get("winrate")])
        for col in "ABCDE":
            ws5.column_dimensions[col].width = 16

    wb.save(XLSX)
    print("wrote", XLSX)
    print("  sheets:", wb.sheetnames)
    print("  Q2 rows:", len(rows))


if __name__ == "__main__":
    sys.exit(main())
